"""Step 2: decide what an email contains.

The model classifies; code decides what happens next. Duplicates are caught by hash before
the model sees anything, and statements and remittances are never routed to posting.
"""

import base64
import csv
import io
import logging
from importlib import resources
from typing import Literal

from anthropic import AnthropicFoundry
from pydantic import BaseModel, Field
from sqlalchemy.orm import Session

from bob import audit
from bob.config import Settings
from bob.models import Attachment, Document, InboundEmail
from bob.storage import Storage

log = logging.getLogger(__name__)

PROMPT_VERSION = "triage_v1"
SYSTEM_PROMPT = resources.files("bob.agent").joinpath("prompts/triage_v1.md").read_text()

DocType = Literal[
    "vendor_invoice",
    "receipt",
    "credit_note",
    "vendor_statement",
    "remittance",
    "gl_export",
    "bank_statement",
    "reply_to_bob",
    "other",
]

# Types that later slices code and post. Everything else is evidence or needs a person.
POSTABLE_TYPES = {"vendor_invoice", "receipt", "credit_note"}
IMAGE_TYPES = {"image/png", "image/jpeg", "image/gif", "image/webp"}
TEXT_TYPES = {"text/plain", "text/csv"}
XLSX_TYPE = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
SPREADSHEET_PREVIEW_ROWS = 200


class TriagedItem(BaseModel):
    attachment_index: int | None = Field(
        description="Index of the attachment this item is, or null for the email body."
    )
    doc_type: DocType
    counterparty: str | None = Field(description="Supplier, customer or person involved.")
    summary: str
    needs_human: bool
    question: str | None


class TriageResult(BaseModel):
    items: list[TriagedItem]


def _spreadsheet_preview(data: bytes) -> str:
    from openpyxl import load_workbook

    wb = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    out = io.StringIO()
    writer = csv.writer(out)
    for ws in wb.worksheets:
        out.write(f"## Sheet: {ws.title}\n")
        for i, row in enumerate(ws.iter_rows(values_only=True)):
            if i >= SPREADSHEET_PREVIEW_ROWS:
                out.write(
                    f"[preview stops after {SPREADSHEET_PREVIEW_ROWS} rows; full file kept]\n"
                )
                break
            writer.writerow(["" if v is None else v for v in row])
    return out.getvalue()


def _attachment_blocks(index: int, att: Attachment, data: bytes) -> list[dict]:
    label = f'Attachment {index}: "{att.filename}" ({att.content_type}, {att.size_bytes} bytes)'
    ctype = att.content_type.lower()
    name = att.filename.lower()
    if ctype == "application/pdf" or name.endswith(".pdf"):
        return [
            {"type": "text", "text": label},
            {
                "type": "document",
                "source": {
                    "type": "base64",
                    "media_type": "application/pdf",
                    "data": base64.b64encode(data).decode(),
                },
                "title": att.filename,
            },
        ]
    if ctype in IMAGE_TYPES:
        return [
            {"type": "text", "text": label},
            {
                "type": "image",
                "source": {
                    "type": "base64",
                    "media_type": ctype,
                    "data": base64.b64encode(data).decode(),
                },
            },
        ]
    if ctype in TEXT_TYPES or name.endswith((".csv", ".txt")):
        text = data.decode("utf-8", errors="replace")
        return [{"type": "text", "text": f"{label}\n<attachment_text>\n{text}\n</attachment_text>"}]
    if ctype == XLSX_TYPE or name.endswith(".xlsx"):
        try:
            preview = _spreadsheet_preview(data)
        except Exception:
            log.exception("Could not read spreadsheet %s", att.filename)
            preview = "[spreadsheet could not be read]"
        return [
            {"type": "text", "text": f"{label}\n<attachment_text>\n{preview}\n</attachment_text>"}
        ]
    return [{"type": "text", "text": f"{label}\n[Bob cannot read this file type.]"}]


def build_messages(
    email: InboundEmail, readable: list[tuple[int, Attachment, bytes]]
) -> list[dict]:
    header = (
        f"From: {email.sender_name or ''} <{email.sender_address}> (sender trust: {email.sender_trust})\n"
        f"Received: {email.received_at.isoformat()}\n"
        f"Subject: {email.subject}\n"
        f"<email_body>\n{email.body_text}\n</email_body>"
    )
    content: list[dict] = [{"type": "text", "text": header}]
    for index, att, data in readable:
        content.extend(_attachment_blocks(index, att, data))
    content.append({"type": "text", "text": "Triage this email."})
    return [{"role": "user", "content": content}]


def classify(
    client: AnthropicFoundry, settings: Settings, messages: list[dict]
) -> TriageResult | None:
    """Returns None if the model declined, so the email is held for a person."""
    response = client.messages.parse(
        model=settings.model,
        max_tokens=16000,
        system=SYSTEM_PROMPT,
        messages=messages,
        thinking={"type": "adaptive"},
        output_config={"effort": settings.triage_effort},
        output_format=TriageResult,
    )
    if response.stop_reason == "refusal":
        log.warning("Triage refused: %s", response.stop_details)
        return None
    if response.stop_reason == "max_tokens" or response.parsed_output is None:
        raise RuntimeError(f"Triage produced no usable output (stop_reason={response.stop_reason})")
    return response.parsed_output


def _status_for(item: TriagedItem) -> str:
    if item.needs_human:
        return "needs_human"
    if item.doc_type in POSTABLE_TYPES:
        return "ready_to_code"
    if item.doc_type in {"vendor_statement", "remittance", "bank_statement"}:
        return "evidence"  # kept for matching and controls, never posted
    if item.doc_type == "gl_export":
        return "migration"  # handled by the one-off import, never posted as activity
    if item.doc_type == "reply_to_bob":
        return "instruction"
    return "no_action"


def triage_email(
    session: Session,
    email_id: int,
    client: AnthropicFoundry,
    storage: Storage,
    settings: Settings,
) -> None:
    email = session.get(InboundEmail, email_id)
    if email is None or email.status != "received":
        return  # already handled (the job may have been retried)

    # Byte-identical duplicates never reach the model.
    readable: list[tuple[int, Attachment, bytes]] = []
    for att in email.attachments:
        if att.duplicate_of_id is not None:
            session.add(
                Document(
                    email=email,
                    attachment_id=att.id,
                    doc_type="duplicate",
                    status="duplicate",
                    summary=f'"{att.filename}" is identical to attachment {att.duplicate_of_id}.',
                )
            )
            continue
        readable.append((len(readable), att, storage.get(att.blob_path)))

    if not readable and not email.body_text.strip():
        email.status = "triaged"
        audit.record(session, "email.triaged", "email", email.id, {"items": 0, "reason": "empty"})
        return

    result = classify(client, settings, build_messages(email, readable))
    if result is None:
        email.status = "held"
        audit.record(session, "email.held", "email", email.id, {"reason": "model_refusal"})
        return

    by_index = {index: att for index, att, _ in readable}
    covered: set[int] = set()
    for item in result.items:
        att = by_index.get(item.attachment_index) if item.attachment_index is not None else None
        if item.attachment_index is not None and att is None:
            log.warning("Triage referenced unknown attachment %s", item.attachment_index)
            continue
        if item.attachment_index is not None:
            covered.add(item.attachment_index)
        doc = Document(
            email=email,
            attachment_id=att.id if att else None,
            doc_type=item.doc_type,
            status=_status_for(item),
            counterparty=item.counterparty,
            summary=item.summary,
            question=item.question,
            model_name=settings.model,
            prompt_version=PROMPT_VERSION,
        )
        # Untrusted senders never feed posting directly: a person confirms first.
        if doc.status == "ready_to_code" and email.sender_trust == "unknown":
            doc.status = "needs_human"
            doc.question = (
                doc.question
                or f"This came from an unknown sender ({email.sender_address}). Is it genuine?"
            )
        session.add(doc)

    # Anything the model skipped still gets a record, so nothing silently disappears.
    for index, att, _ in readable:
        if index not in covered:
            session.add(
                Document(
                    email=email,
                    attachment_id=att.id,
                    doc_type="other",
                    status="needs_human",
                    summary=f'"{att.filename}" was not classified.',
                    question=f'What is "{att.filename}"?',
                    model_name=settings.model,
                    prompt_version=PROMPT_VERSION,
                )
            )

    email.status = "triaged"
    session.flush()
    audit.record(
        session,
        "email.triaged",
        "email",
        email.id,
        {
            "model": settings.model,
            "prompt_version": PROMPT_VERSION,
            "documents": [
                {"id": d.id, "type": d.doc_type, "status": d.status, "summary": d.summary}
                for d in email.documents
            ],
        },
    )
